import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
import os

def generate_excel_report(test_cases: list, url: str) -> str:
    """
    Generate Excel report from test cases
    Returns the file path
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    # ── Colors ──
    header_fill    = PatternFill("solid", fgColor="4B0082")
    pass_fill      = PatternFill("solid", fgColor="C6EFCE")
    fail_fill      = PatternFill("solid", fgColor="FFC7CE")
    pending_fill   = PatternFill("solid", fgColor="FFEB9C")
    title_fill     = PatternFill("solid", fgColor="1a1a2e")

    # ── Title Row ──
    ws.merge_cells("A1:G1")
    title_cell = ws["A1"]
    title_cell.value = f"QA Pilot — Test Report | {url} | {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    title_cell.font = Font(bold=True, color="FFFFFF", size=13)
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # ── Headers ──
    headers = ["Test ID", "Description", "Steps", "Expected Result", "Actual Result", "Status", "Remarks"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )
    ws.row_dimensions[2].height = 25

    # ── Column Widths ──
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 30
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 25

    # ── Data Rows ──
    for row_idx, tc in enumerate(test_cases, 3):
        status = tc.get("status", "Pending").strip()

        if status.lower() == "pass":
            status_fill = pass_fill
            status_color = "375623"
        elif status.lower() == "fail":
            status_fill = fail_fill
            status_color = "9C0006"
        else:
            status_fill = pending_fill
            status_color = "7D6608"

        row_data = [
            tc.get("test_id", f"TC_{row_idx-2:03d}"),
            tc.get("description", ""),
            tc.get("steps", ""),
            tc.get("expected", ""),
            tc.get("actual", ""),
            status,
            tc.get("remarks", "")
        ]

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin")
            )
            if col_idx == 6:  # Status column
                cell.fill = status_fill
                cell.font = Font(bold=True, color=status_color)
                cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[row_idx].height = 60

    # ── Summary Sheet ──
    ws2 = wb.create_sheet("Summary")
    ws2.column_dimensions["A"].width = 25
    ws2.column_dimensions["B"].width = 20

    summary_title = ws2.cell(row=1, column=1, value="Test Execution Summary")
    summary_title.font = Font(bold=True, size=14, color="FFFFFF")
    summary_title.fill = title_fill
    ws2.merge_cells("A1:B1")
    summary_title.alignment = Alignment(horizontal="center")
    ws2.row_dimensions[1].height = 30

    total     = len(test_cases)
    passed    = sum(1 for t in test_cases if t.get("status","").lower() == "pass")
    failed    = sum(1 for t in test_cases if t.get("status","").lower() == "fail")
    pending   = total - passed - failed
    pass_rate = f"{(passed/total*100):.1f}%" if total > 0 else "0%"

    summary_data = [
        ("Application URL", url),
        ("Test Date", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Total Test Cases", total),
        ("Passed", passed),
        ("Failed", failed),
        ("Pending", pending),
        ("Pass Rate", pass_rate),
    ]

    for i, (label, value) in enumerate(summary_data, 2):
        label_cell = ws2.cell(row=i, column=1, value=label)
        label_cell.font = Font(bold=True)
        label_cell.fill = PatternFill("solid", fgColor="E8E0F0")
        label_cell.border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        value_cell = ws2.cell(row=i, column=2, value=value)
        value_cell.border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        ws2.row_dimensions[i].height = 22

    # ── Save File ──
    os.makedirs("reports/output", exist_ok=True)
    filename = f"reports/output/qa_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(filename)
    return filename